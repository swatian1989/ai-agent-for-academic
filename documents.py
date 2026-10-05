"""
documents.py — getting text in and documents out.

    read_upload(name, data)   .md / .txt / .pdf / .docx  ─► plain text (for notes, manuscripts, references)
    markdown_to_docx(md)      Markdown ─► Word (.docx) bytes: headings, lists, tables, bold/italic
    papers_to_bibtex(papers)  ─► BibTeX for Zotero / EndNote / Mendeley
    papers_to_ris(papers)     ─► RIS for Zotero / EndNote / Mendeley
"""
from __future__ import annotations

import io
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# ------------------------------------------------------------------------- uploads


def read_upload(name: str, data: bytes) -> str:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        lines = []
        for p in doc.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            text = p.text.strip()
            if not text:
                continue
            m = re.match(r"heading (\d)", style)
            lines.append(("#" * int(m.group(1)) + " " + text) if m else
                         ("# " + text) if style == "title" else text)
        for t in doc.tables:
            for i, row in enumerate(t.rows):
                lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
                if i == 0:
                    lines.append("|" + "---|" * len(row.cells))
        return "\n\n".join(lines)
    return data.decode("utf-8", errors="replace")


# ------------------------------------------------------------------- Markdown → Word
_INLINE = re.compile(r"(\*\*[^*]+?\*\*|__[^_]+?__|\*[^*\s][^*]*?\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\))")


def _add_inline(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            _add_inline(paragraph, part[2:-2], bold=True, italic=italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _add_inline(paragraph, part[1:-1], bold=bold, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            run = paragraph.add_run(f"{label} ({url})")
            run.bold, run.italic = bold, italic
        else:
            run = paragraph.add_run(part)
            run.bold, run.italic = bold, italic


def _is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", line))


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def markdown_to_docx(md: str, title: str | None = None) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    if title:
        doc.core_properties.title = title

    lines = md.replace("\r\n", "\n").split("\n")
    i, para = 0, []

    def flush():
        if para:
            p = doc.add_paragraph()
            _add_inline(p, " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):                                    # code block
            flush()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                run = doc.add_paragraph().add_run(lines[i])
                run.font.name, run.font.size = "Courier New", Pt(9)
                i += 1
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", stripped)
        if m:                                                             # heading
            flush()
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 4))
            _add_inline(h, m.group(2))
            if level == 1 and i < 3:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _is_table_sep(lines[i + 1]):   # table
            flush()
            header = _cells(stripped)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.paragraphs[0].text = ""
                _add_inline(cell.paragraphs[0], text, bold=True)
            for r, row in enumerate(rows, start=1):
                for c in range(len(header)):
                    _add_inline(table.rows[r].cells[c].paragraphs[0], row[c] if c < len(row) else "")
            doc.add_paragraph()
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):                  # horizontal rule
            flush()
            i += 1
            continue
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if m:                                                             # list item
            flush()
            numbered = m.group(2)[0].isdigit()
            depth = len(m.group(1).replace("\t", "    ")) // 2
            style = ("List Number" if numbered else "List Bullet") + (" 2" if depth else "")
            _add_inline(doc.add_paragraph(style=style), m.group(3))
            i += 1
            continue
        if stripped.startswith(">"):                                      # block quote
            flush()
            p = doc.add_paragraph(style="Quote")
            _add_inline(p, stripped.lstrip("> ").strip())
            i += 1
            continue
        if not stripped:
            flush()
        else:
            para.append(stripped)
        i += 1
    flush()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- references
def _ascii(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()


def _doi(p) -> str:
    if p.id.startswith("doi:"):
        return p.id[4:]
    m = re.search(r"doi\.org/(10\.[^\s]+)", p.url or "")
    return m.group(1) if m else ""


def papers_to_bibtex(papers) -> str:
    entries, used = [], set()
    for p in papers:
        last = re.sub(r"[^A-Za-z]", "", _ascii(p.authors[0].split()[-1])) if p.authors else "anon"
        word = next((re.sub(r"[^A-Za-z]", "", _ascii(w)) for w in p.title.split() if len(w) > 3), "paper")
        key = f"{last.lower()}{p.year or ''}{word.lower()}"
        k, n = key, 2
        while k in used:
            k, n = f"{key}{chr(96 + n)}", n + 1
        used.add(k)
        esc = lambda s: str(s).replace("{", "\\{").replace("}", "\\}")  # noqa: E731
        fields = {"title": esc(p.title), "author": " and ".join(esc(a) for a in p.authors),
                  "year": p.year or "", "journal": esc(p.venue), "doi": _doi(p), "url": p.url,
                  "note": f"Source id: {p.id}"}
        body = ",\n".join(f"  {f} = {{{v}}}" for f, v in fields.items() if v)
        entries.append(f"@article{{{k},\n{body}\n}}")
    return "\n\n".join(entries) + "\n"


def papers_to_ris(papers) -> str:
    out = []
    for p in papers:
        lines = ["TY  - JOUR", f"TI  - {p.title}"]
        lines += [f"AU  - {a}" for a in p.authors]
        if p.year:
            lines.append(f"PY  - {p.year}")
        if p.venue:
            lines.append(f"JO  - {p.venue}")
        if _doi(p):
            lines.append(f"DO  - {_doi(p)}")
        if p.url:
            lines.append(f"UR  - {p.url}")
        if p.abstract:
            lines.append(f"AB  - {p.abstract}")
        lines += [f"ID  - {p.id}", "ER  - "]
        out.append("\n".join(lines))
    return "\n\n".join(out) + "\n"
